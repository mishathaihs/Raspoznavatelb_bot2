from __future__ import annotations

import asyncio
import io
from pathlib import Path

import pdfplumber
from docx import Document

from bot.models import ParsedDocument
from services.gpt_vision import get_gpt_client

SAMPLES = [
    "tests/samples/nakladnaya.pdf",
    "tests/samples/ttn.pdf",
    "tests/samples/akt_schet.pdf",
    "tests/samples/schet_protokol.pdf",
    "tests/samples/nakladnaya2.pdf",
]


def _extract_docx_text(path: Path) -> str:
    doc = Document(path)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            lines.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_first_page_image(path: Path) -> bytes:
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            raise ValueError("PDF has no pages")
        page = pdf.pages[0]
        image = page.to_image(resolution=250)
        buffer = io.BytesIO()
        image.original.save(buffer, format="PNG")
        return buffer.getvalue()


async def _process_file(client, path: Path) -> None:
    image_bytes = None
    text_from_docx = None
    if path.suffix.lower() == ".docx":
        text_from_docx = _extract_docx_text(path)
    elif path.suffix.lower() == ".pdf":
        image_bytes = _extract_first_page_image(path)
    else:
        image_bytes = path.read_bytes()

    parsed: ParsedDocument = await client.parse_document(
        image_bytes=image_bytes, text_from_docx=text_from_docx, filename=path.name
    )
    print(
        f"File: {path.name}\n"
        f"  doc_type={parsed.doc_type}\n  number={parsed.number}\n  date={parsed.date}\n"
        f"  main_counterparty={parsed.main_counterparty}\n  total={getattr(parsed.totals, 'amount_with_vat', None)}\n"
        f"  warnings={parsed.warnings}\n"
    )


async def main() -> None:
    client = get_gpt_client()
    for sample in SAMPLES:
        path = Path(sample)
        if not path.exists():
            print(f"Skip missing sample: {path}")
            continue
        try:
            await _process_file(client, path)
        except Exception as exc:  # pragma: no cover - debug helper
            print(f"Failed to parse {path}: {exc}")


if __name__ == "__main__":
    asyncio.run(main())
