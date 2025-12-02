from datetime import date

from docx import Document

from bot.models import ParsedDocument, Party
from services import docx_builder


def test_sanitize_replacements_converts_none_and_numbers():
    raw = {"{{A}}": None, "{{B}}": 5, "{{C}}": 2.5, "{{D}}": "text"}

    sanitized = docx_builder._sanitize_replacements(raw)

    assert all(isinstance(v, str) for v in sanitized.values())
    assert sanitized["{{A}}"] == ""
    assert sanitized["{{B}}"] == "5"
    assert sanitized["{{C}}"] == "2.5"
    assert sanitized["{{D}}"] == "text"


def test_render_docx_handles_none_replacements(monkeypatch, tmp_path):
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "out.docx"

    template = Document()
    template.add_paragraph("Doc {{DOC_NUMBER}} basis {{BASIS}}").alignment = 0
    template.save(template_path)

    monkeypatch.setattr(
        docx_builder,
        "_build_replacements",
        lambda parsed: {"{{DOC_NUMBER}}": "123", "{{BASIS}}": None},
    )

    parsed = ParsedDocument(doc_type="акт", number="123")
    docx_builder.render_docx(parsed, template_path, output_path)

    result = Document(output_path)
    assert "{{" not in " ".join(p.text for p in result.paragraphs)


def test_render_docx_with_missing_fields_creates_file(tmp_path):
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "out.docx"

    template = Document()
    template.add_paragraph("Номер {{DOC_NUMBER}} от {{DOC_DATE}} для {{CUSTOMER}}").alignment = 0
    template.save(template_path)

    parsed = ParsedDocument(
        doc_type="акт",
        number=None,
        date=date(2024, 1, 10),
        buyer=Party(name=None),
    )

    docx_builder.render_docx(parsed, template_path, output_path)

    result = Document(output_path)
    text = " ".join(p.text for p in result.paragraphs)
    assert "{{" not in text
    assert "10.01.2024" in text
