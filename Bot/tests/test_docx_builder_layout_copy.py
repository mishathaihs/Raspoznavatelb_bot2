from datetime import date
from decimal import Decimal

from docx import Document

from bot.models import Item, ParsedDocument, Party, Totals
from services.docx_builder import render_docx


def _collect_text(doc: Document) -> str:
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


def test_layout_copy_uses_template_slots(tmp_path):
    template_path = tmp_path / "template.docx"
    template = Document()
    template.add_paragraph("Номер {{DOC_NUMBER}} от {{DOC_DATE}} для {{CUSTOMER}}").alignment = 0
    table = template.add_table(rows=2, cols=5)
    headers = ["Наименование", "Сумма без НДС", "НДС", "Итого", "Валюта"]
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    template_row = table.rows[1]
    template_row.cells[0].text = "{{POSITION_NAME}}"
    template_row.cells[1].text = "{{POSITION_AMOUNT_WO_VAT}}"
    template_row.cells[2].text = "{{POSITION_VAT}}"
    template_row.cells[3].text = "{{POSITION_WITH_VAT}}"
    template_row.cells[4].text = "{{CURRENCY}}"
    template.save(template_path)

    items = [
        Item(name="Услуга 1", amount_without_vat=Decimal("10"), vat_amount=Decimal("2"), amount_with_vat=Decimal("12")),
        Item(name="Услуга 2", amount_without_vat=Decimal("5"), vat_amount=Decimal("1"), amount_with_vat=Decimal("6")),
    ]

    parsed = ParsedDocument(
        doc_type="акт",
        number="123",
        date=date(2024, 1, 10),
        buyer=Party(name='ООО "Клиент"'),
        executor='ООО "Исполнитель"',
        items=items,
        totals=Totals(amount_with_vat=Decimal("18"), vat_amount=Decimal("3"), amount_without_vat=Decimal("15")),
    )

    output_path = tmp_path / "result.docx"
    render_docx(parsed, template_path, output_path, layout_copy=True)

    result = Document(output_path)
    table = result.tables[0]

    assert len(table.rows) == 1 + len(items)
    replaced_text = _collect_text(result)
    assert "{{" not in replaced_text
    assert "Услуга 1" in table.rows[1].cells[0].text
    assert "Услуга 2" in table.rows[2].cells[0].text

