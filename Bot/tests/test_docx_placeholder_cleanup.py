from pathlib import Path

import pytest
from docx import Document

from services import docx_builder


def test_unreplaced_placeholders_are_removed(tmp_path: Path, caplog):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("Итого без НДС: {{ TOTAL_WITHOUT_VAT }} {{ CURRENCY }}")
    doc.save(template)

    with caplog.at_level("WARNING"):
        # Provide replacements without the placeholders above to force cleanup
        doc_template = Document(template)
        docx_builder._replace_in_document(doc_template, {"{{OTHER}}": "value"})
        doc_template.save(output)

    rendered = Document(output)
    text = "\n".join(p.text for p in rendered.paragraphs)
    assert "{{" not in text and "}}" not in text
    assert any("Unreplaced placeholder" in message for message in caplog.messages)
