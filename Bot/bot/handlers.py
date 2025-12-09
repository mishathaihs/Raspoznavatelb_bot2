from __future__ import annotations

import asyncio
import io
import logging
import re
import tempfile
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pdfplumber
from aiogram import F, Router
import html

from aiogram.enums import ContentType
from aiogram.exceptions import TelegramBadRequest
from aiogram.filters import Command, CommandObject
from aiogram.types import CallbackQuery, InlineKeyboardMarkup, Message
from aiogram.utils.keyboard import InlineKeyboardBuilder
from docx import Document

from bot.models import Item, ParsedDocument, Party, Totals
from config import get_settings
from services import get_drive_service, get_sheets_service, get_vision_extractor
from services.layout_docx_builder import build_layout_docx
from services.layout_types import LayoutExtraction
from services.docx_builder import ensure_default_templates, render_docx
from services.company_resolver import resolve_parties
from services.exceptions import (
    DriveQuotaExceededError,
    DriveUploadError,
    ParseError,
    PersistDriveError,
    SpreadsheetInitError,
)
from services.parsing_utils import detect_currency
from services.amounts import prefer_text_amount
from services.formatting import format_money_ru
from services.normalization import strip_address
from services.vision_extractor import VisionExtractionError

logger = logging.getLogger(__name__)

settings = get_settings()

TELEGRAM_DOWNLOAD_TIMEOUT = 300

router = Router()


def esc(text: str | None) -> str:
    return html.escape(text or "")


def _vision_error_message(exc: VisionExtractionError) -> str:
    message_text = str(exc)
    lowered = message_text.lower()
    if "json" in lowered:
        return (
            "Не удалось обработать документ из-за технической ошибки при разборе ответа модели. "
            "Это не связано с качеством фото. Попробуйте отправить документ ещё раз чуть позже или другой документ."
        )
    if "контрагент" in lowered:
        return "Не удалось надёжно распознать названия контрагентов в документе, проверьте вручную."
    return (
        "Не удалось распознать документ из-за технической ошибки. Попробуйте ещё раз или отправьте более чёткое фото."
    )


@dataclass
class PendingUpload:
    internal_doc_id: str
    doc_data: ParsedDocument
    file_bytes: bytes
    suffix: str
    user_id: int | None
    layout: LayoutExtraction | None = None
    duplicate_order_number: int | None = None
    duplicate_file_id: str | None = None


PENDING_UPLOADS: dict[str, PendingUpload] = {}


START_TEXT = (
    "👋 Привет! Я помогу оформить акт или накладную.\n\n"
    "Пошаговая инструкция:\n"
    "1️⃣ Шаг 1 — пришлите файл PDF/JPG/PNG с актом или накладной.\n"
    "2️⃣ Шаг 2 — я распознаю реквизиты, внесу строки в Google Sheets и переименую файл.\n"
    "3️⃣ Шаг 3 — в ответ вы получите сводку и кнопку для удаления при необходимости.\n\n"
    "🔍 Чтобы бот корректно распознал документ, подготовьте файл:\n"
    "• Формат: PDF / JPG / PNG\n"
    "• Разрешение: от 200–300 dpi, для фото — ширина 1500–2000 px\n"
    "• Документ целиком в кадре, без сильного наклона и бликов.\n\n"
    "Готовы? Просто отправьте документ, и я подскажу, что делать дальше!"
)

HELP_TEXT = (
    "ℹ️ Памятка:\n"
    "• Отправьте один файл с актом или накладной (PDF/PNG/JPG).\n"
    "• Дождитесь сообщения со сводкой — там будет номер по порядку и кнопка для удаления.\n"
    "• Чтобы удалить документ позже, используйте команду /delete «номер» или кнопку 🗑️ под сообщением.\n\n"
    "🔍 Требования к качеству:\n"
    "• Хорошее освещение, документ без бликов и сильного наклона.\n"
    "• Текст не меньше 10–12 pt; для фото ширина 1500–2000 px.\n"
    "• Для сканов и PDF используйте разрешение от 200–300 dpi."
)

DELETE_HELP_TEXT = (
    "🗑️ Как удалить документ:\n"
    "1. Найдите сообщение со сводкой и нажмите кнопку ‘🗑️ Удалить документ’.\n"
    "2. Либо отправьте команду /delete «номер по порядку».\n"
    "После удаления можно отправить исправленный файл повторно."
)

SPREADSHEET_ERROR_MESSAGE = (
    "⚠️ Не удалось записать документ в Google Sheets. Проверьте GOOGLE_SHEETS_SPREADSHEET_ID "
    "и доступ сервисного аккаунта, затем перезапустите бота."
)

DRIVE_QUOTA_ERROR_MESSAGE = (
    "Не удалось сохранить документ в Google Диск: хранилище недоступно или закончилась квота.\n"
    "Распознать акт/накладную получилось, но файл не был прикреплён.\n"
    "Обратитесь к администратору, чтобы освободить место или перенастроить доступ."
)


TEMPLATE_MAP = {
    "акт": settings.docx_act_template,
    "товарная_накладная": settings.docx_nakladnaya_template,
    "товарно-транспортная_накладная": settings.docx_ttn_template,
    "счет-фактура": settings.docx_schet_faktura_template,
    "счет-протокол": settings.docx_schet_protokol_template,
}
DEFAULT_TEMPLATE = settings.docx_other_template


def build_processed_document_keyboard(order_number: int) -> InlineKeyboardMarkup:
    builder = InlineKeyboardBuilder()
    builder.button(text=f"🗑️ Удалить документ №{order_number}", callback_data=f"delete:{order_number}")
    builder.button(text="📤 Обработать ещё документ", callback_data="process:more")
    builder.adjust(1)
    return builder.as_markup()


def build_start_keyboard() -> InlineKeyboardMarkup:
    builder = InlineKeyboardBuilder()
    builder.button(text="📖 Памятка", callback_data="help:guide")
    builder.button(text="🗑️ Как удалить документ", callback_data="help:delete")
    return builder.as_markup()


@router.message(Command("start"))
async def handle_start(message: Message) -> None:
    await message.answer(START_TEXT, reply_markup=build_start_keyboard())


@router.message(Command("help"))
async def handle_help(message: Message) -> None:
    await message.answer(HELP_TEXT)


@router.callback_query(F.data == "help:guide")
async def handle_help_callback(query: CallbackQuery) -> None:
    await query.answer()
    await query.message.answer(HELP_TEXT)


@router.callback_query(F.data == "help:delete")
async def handle_delete_help_callback(query: CallbackQuery) -> None:
    await query.answer()
    await query.message.answer(DELETE_HELP_TEXT)


@router.message(Command("delete"))
async def handle_delete_command(message: Message, command: CommandObject) -> None:
    if not command.args or not command.args.isdigit():
        await message.answer(
            "Пожалуйста, укажите номер по порядку: /delete 10\n\n" + DELETE_HELP_TEXT
        )
        return
    order_number = int(command.args)
    await _delete_document(order_number, message)


@router.callback_query(F.data.startswith("delete:"))
async def handle_delete_callback(query: CallbackQuery) -> None:
    await query.answer()
    try:
        order_number = int(query.data.split(":", 1)[1])
    except (IndexError, ValueError):
        await query.message.answer("Некорректный запрос на удаление.")
        return
    await _delete_document(order_number, query.message)


@router.callback_query(F.data == "process:more")
async def handle_process_more_callback(query: CallbackQuery) -> None:
    await query.answer()
    logger.info(
        "User %s requested to process another document",
        getattr(query.from_user, "id", "unknown"),
    )
    await query.message.answer(
        "📤 Отправьте следующий акт или накладную файлом — я обработаю его так же, как предыдущий."
    )


@router.callback_query(F.data.startswith("dup:"))
async def handle_duplicate_resolution(query: CallbackQuery) -> None:
    await query.answer()
    try:
        _, action, pending_id = query.data.split(":", 2)
    except ValueError:
        await query.message.answer("Некорректный выбор действия.")
        return

    pending = PENDING_UPLOADS.pop(pending_id, None)
    if pending is None:
        await query.message.answer("Данные документа не найдены. Отправьте файл ещё раз.")
        return

    if action == "keep":
        await query.message.answer("Оставил предыдущую запись без изменений. Новый файл не сохранён.")
        return

    if action != "replace":
        await query.message.answer("Некорректное действие. Отправьте файл заново.")
        return

    sheets = get_sheets_service()
    drive = get_drive_service()
    if pending.duplicate_order_number:
        file_id = await sheets.delete_document_rows_by_order_number(pending.duplicate_order_number)
        if file_id:
            try:
                await drive.delete_file(file_id)
            except Exception:  # pragma: no cover - remote API
                logger.exception("Failed to delete old Drive file %s during replace", file_id)
    await _persist_document(pending, query.message)


async def _delete_document(order_number: int, message: Message) -> None:
    sheets = get_sheets_service()
    drive = get_drive_service()
    logger.info("Received request to delete document order=%s", order_number)
    try:
        file_id = await sheets.delete_document_rows_by_order_number(order_number)
    except Exception as exc:  # pragma: no cover - depends on remote API
        logger.exception("Failed to delete rows for order %s: %s", order_number, exc)
        await message.answer("Не удалось удалить строки в Google Sheets. Попробуйте чуть позже.")
        return

    if file_id is None:
        logger.info("Document order=%s not found in Sheets", order_number)
        await message.answer(
            f"Записи с номером {order_number} не найдены. Проверьте номер и попробуйте снова.\n\n"
            + DELETE_HELP_TEXT
        )
        return

    drive_error = None
    try:
        await drive.delete_file(file_id or "")
    except Exception as exc:  # pragma: no cover - depends on remote API
        drive_error = exc
        logger.exception("Failed to delete Drive file %s: %s", file_id, exc)

    if drive_error:
        await message.answer(
            "Строки удалены из таблицы, но файл на Google Диске удалить не удалось. Попробуйте позже или удалите вручную."
        )
    else:
        logger.info(
            "Deleted document order=%s internal_doc_id=%s from Sheets and Drive",
            order_number,
            "n/a",
        )
        await message.answer(
            f"Документ №{order_number} удалён из таблицы и Google Диска. Можете отправить исправленный файл в любой момент."
        )


async def _persist_document(pending: PendingUpload, message: Message) -> None:
    sheets = get_sheets_service()
    drive = get_drive_service()

    reservation = None
    try:
        reservation = await sheets.get_next_order_number()
        order_number = reservation.order_number
    except SpreadsheetInitError as exc:
        logger.exception("Failed to initialize Sheets while reserving order number: %s", exc)
        await message.answer(SPREADSHEET_ERROR_MESSAGE)
        return
    except Exception as exc:  # pragma: no cover - depends on remote API
        logger.exception("Failed to reserve order number: %s", exc)
        await message.answer("Не удалось получить следующий номер по порядку. Попробуйте повторить позже.")
        return

    logger.info("[doc_id=%s] Reserved order_number=%s", pending.internal_doc_id, order_number)

    tmp_path: Path | None = None
    docx_path: Path | None = None
    layout_docx_path: Path | None = None
    file_id: str | None = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=pending.suffix) as tmp_file:
            tmp_file.write(pending.file_bytes)
            tmp_path = Path(tmp_file.name)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as structured:
            docx_path = Path(structured.name)

        ensure_default_templates()

        template_path = Path(TEMPLATE_MAP.get(pending.doc_data.doc_type, DEFAULT_TEMPLATE))
        if not template_path.is_absolute():
            template_path = PROJECT_ROOT / template_path
        if not template_path.exists():
            raise FileNotFoundError(f"Docx template not found: {template_path}")
        try:
            totals = pending.doc_data.totals or Totals()
            resolved_parties = resolve_parties(pending.doc_data.supplier_raw, pending.doc_data.customer_raw)
            render_docx(
                pending.doc_data,
                template_path,
                docx_path,
                executor_name=resolved_parties.executor_name
                if resolved_parties.executor_in_dict
                else (resolved_parties.executor_raw or pending.doc_data.supplier_raw),
                buyer_name=resolved_parties.buyer_name
                if resolved_parties.buyer_in_dict
                else (resolved_parties.buyer_raw or pending.doc_data.customer_raw),
                doc_number=pending.doc_data.number or pending.doc_data.doc_number,
                doc_date=pending.doc_data.date or pending.doc_data.doc_date,
                totals_without_vat=totals.total_without_vat,
                totals_vat=totals.vat_amount,
                totals_with_vat=totals.total_with_vat,
                currency=pending.doc_data.currency or pending.doc_data.currency_code,
            )
        except Exception as exc:
            logger.exception(
                "[order=%s doc_id=%s] Failed to render docx with template %s: %s",
                order_number,
                pending.internal_doc_id,
                template_path,
                exc,
            )
            raise

        layout_copy_uploaded = False
        layout_file_id: str | None = None
        if settings.enable_layout_docx_copy and pending.layout and pending.layout.words:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as layout_tmp:
                layout_docx_path = Path(layout_tmp.name)
            try:
                build_layout_docx(
                    pending.layout.words,
                    pending.layout.effective_text,
                    layout_docx_path,
                    title=_build_layout_title(pending.doc_data),
                )
            except Exception:
                logger.exception(
                    "[order=%s doc_id=%s] Failed to build layout docx copy",
                    order_number,
                    pending.internal_doc_id,
                )

        totals = pending.doc_data.totals or Totals()
        total_with_vat = totals.total_with_vat or totals.amount_with_vat
        logger.info(
            "[order=%s doc_id=%s] Parsed doc_type=%s total_with_vat=%s layout_copy=%s",
            order_number,
            pending.internal_doc_id,
            pending.doc_data.doc_type,
            total_with_vat,
            bool(layout_docx_path),
        )

        docx_filename = _build_drive_filename(order_number, pending.doc_data, ".docx")
        try:
            file_id = await drive.upload_generated(docx_path, docx_filename)
            logger.info(
                "[order=%s doc_id=%s] Structured docx uploaded to Drive file_id=%s",
                order_number,
                pending.internal_doc_id,
                file_id,
            )
            if layout_docx_path and layout_docx_path.exists():
                copy_filename = _build_layout_copy_filename(order_number, pending.doc_data, ".docx")
                try:
                    layout_file_id = await drive.upload_generated(layout_docx_path, copy_filename)
                    layout_copy_uploaded = True
                    logger.info(
                        "[order=%s doc_id=%s] Layout docx copy uploaded file_id=%s",
                        order_number,
                        pending.internal_doc_id,
                        layout_file_id,
                    )
                except Exception:
                    logger.exception(
                        "[order=%s doc_id=%s] Failed to upload layout copy",
                        order_number,
                        pending.internal_doc_id,
                    )
            original_filename = _build_drive_filename(order_number, pending.doc_data, pending.suffix)
            try:
                await drive.upload_original(tmp_path, original_filename)
            except Exception:
                logger.warning(
                    "[order=%s doc_id=%s] Failed to upload original file, docx uploaded",
                    order_number,
                    pending.internal_doc_id,
                )
        except DriveQuotaExceededError as exc:
            logger.error(
                "[doc_id=%s order=%s] Drive quota exceeded while uploading %s: %s",
                pending.internal_doc_id,
                order_number,
                docx_filename,
                exc,
            )
            await sheets.rollback_order_reservation(reservation)
            raise PersistDriveError("drive_quota") from exc
        except DriveUploadError as exc:
            logger.exception(
                "[doc_id=%s order=%s] Drive upload failed for %s: %s",
                pending.internal_doc_id,
                order_number,
                docx_filename,
                exc,
            )
            if not settings.allow_sheets_without_drive_file:
                await sheets.rollback_order_reservation(reservation)
                raise PersistDriveError("drive_upload") from exc
            logger.warning(
                "[doc_id=%s order=%s] Continuing without Drive file due to allow_sheets_without_drive_file",
                pending.internal_doc_id,
                order_number,
            )
            file_id = None

        await sheets.append_document_rows(pending.doc_data, pending.internal_doc_id, order_number, file_id=file_id)
        logger.info(
            "[order=%s doc_id=%s] Appended rows to Sheets",
            order_number,
            pending.internal_doc_id,
        )

        summary_text = _build_summary(pending.doc_data, order_number, file_saved=bool(file_id))
        try:
            await message.answer(
                summary_text,
                reply_markup=build_processed_document_keyboard(order_number),
                parse_mode="HTML",
            )
        except TelegramBadRequest as exc:
            logger.exception("Telegram returned error while sending summary: %s", exc)
            if "can't parse entities" in str(exc):
                await message.answer(
                    summary_text,
                    reply_markup=build_processed_document_keyboard(order_number),
                    parse_mode=None,
                    disable_web_page_preview=True,
                )
            else:
                raise
    except PersistDriveError:
        raise
    except SpreadsheetInitError as exc:
        logger.exception("Failed to initialize Sheets while saving order=%s: %s", order_number, exc)
        await message.answer(SPREADSHEET_ERROR_MESSAGE)
        try:
            if file_id:
                await drive.delete_file(file_id)
        except Exception:
            logger.exception("Failed to cleanup Drive file after Sheets error")
        await sheets.rollback_order_reservation(reservation)
    except Exception as exc:  # pragma: no cover - remote API
        logger.exception("Failed to persist document order=%s: %s", order_number, exc)
        await message.answer("Не удалось сохранить документ. Попробуйте ещё раз позже.")
        try:
            if file_id:
                await drive.delete_file(file_id)
        except Exception:
            logger.exception("Failed to cleanup Drive file after error")
        await sheets.rollback_order_reservation(reservation)
    finally:
        for path in (tmp_path, docx_path, layout_docx_path):
            if path:
                try:
                    path.unlink(missing_ok=True)
                except PermissionError:
                    logger.warning("Could not delete temp file %s (probably locked), will be cleaned later", path)


@router.message(F.content_type.in_({ContentType.DOCUMENT, ContentType.PHOTO}))
async def handle_document(message: Message) -> None:
    document = message.document or (message.photo[-1] if message.photo else None)
    if document is None:
        await message.answer("Файл не распознан. Пожалуйста, отправьте документ повторно.")
        return

    if message.document and message.document.mime_type not in {
        "application/pdf",
        "image/png",
        "image/jpeg",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }:
        await message.answer("Поддерживаются только PDF, PNG, JPEG или DOCX.")
        return

    internal_doc_id = str(uuid.uuid4())
    user_id = getattr(message.from_user, "id", None)
    doc_data: ParsedDocument | None = None
    pending: PendingUpload | None = None
    duplicate_info = None

    logger.info(
        "Received document for processing from user=%s internal_doc_id=%s filename=%s content_type=%s mime=%s",
        user_id,
        internal_doc_id,
        getattr(document, "file_name", None),
        message.content_type,
        getattr(document, "mime_type", None),
    )

    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            original_name = getattr(document, "file_name", None)
            if original_name:
                suffix = Path(original_name).suffix or ".pdf"
                filename = original_name
            else:
                suffix = ".jpg" if message.photo else ".pdf"
                filename = f"document_{uuid.uuid4()}{suffix}"
            tmp_path = Path(tmp_dir) / filename
            await message.bot.download(
                document, destination=tmp_path, timeout=TELEGRAM_DOWNLOAD_TIMEOUT
            )

            image_bytes: bytes | None = None
            text_from_docx: str | None = None
            if tmp_path.suffix.lower() == ".pdf":
                image_bytes = _extract_first_page_image(tmp_path)
            elif tmp_path.suffix.lower() in {".jpg", ".jpeg", ".png"}:
                image_bytes = tmp_path.read_bytes()
            elif tmp_path.suffix.lower() == ".docx":
                text_from_docx = _extract_docx_text(tmp_path)

            vision = get_vision_extractor()

            async def _run_vision(strict: bool = False) -> ParsedDocument:
                return await vision.extract(
                    image_bytes=image_bytes,
                    text_content=text_from_docx,
                    strict_json=strict,
                )

            try:
                doc_data = await _run_vision()
            except VisionExtractionError as exc:
                if "некорректный json" in str(exc).lower() or "not valid json" in str(exc).lower():
                    logger.warning(
                        "[doc_id=%s user_id=%s] Vision JSON parse failed, retrying with strict prompt",
                        internal_doc_id,
                        user_id,
                    )
                    doc_data = await _run_vision(strict=True)
                else:
                    raise

            layout_data = vision.get_last_layout()
            _postprocess_document(doc_data, settings.our_org_name)

            vision_model = getattr(vision, "_model", settings.openai_vision_model)
            logger.info(
                "[doc_id=%s] Vision model=%s layout_words=%s parsed_items=%s totals_with_vat=%s",
                internal_doc_id,
                vision_model,
                len(layout_data.words) if layout_data else 0,
                len(doc_data.items),
                getattr(doc_data.totals, "total_with_vat", None) if doc_data.totals else None,
            )

            file_bytes = tmp_path.read_bytes()
            pending = PendingUpload(
                internal_doc_id=internal_doc_id,
                doc_data=doc_data,
                file_bytes=file_bytes,
                suffix=tmp_path.suffix,
                user_id=user_id,
                layout=layout_data,
            )

        sheets = get_sheets_service()
        await sheets.ensure_initialized()
        duplicate_info = await sheets.find_duplicate(doc_data)
    except asyncio.TimeoutError:
        logger.error(
            "[doc_id=%s user_id=%s] Timeout while downloading file from Telegram",
            internal_doc_id,
            user_id,
        )
        await message.answer(
            "Не удалось скачать документ из Telegram из-за медленного соединения. "
            "Попробуйте ещё раз отправить файл (как документ) или повторите позже."
        )
        return
    except VisionExtractionError as exc:
        logger.exception(
            "[doc_id=%s user_id=%s] Vision parsing failed for file name=%s mime=%s: %s",
            internal_doc_id,
            user_id,
            getattr(document, "file_name", None),
            getattr(document, "mime_type", None),
            exc,
        )
        await message.answer(_vision_error_message(exc))
        return
    except ParseError as exc:
        logger.exception(
            "[doc_id=%s user_id=%s] Vision parsing failed for file name=%s mime=%s: %s",
            internal_doc_id,
            user_id,
            getattr(document, "file_name", None),
            getattr(document, "mime_type", None),
            exc,
        )
        await message.answer(
            "Не удалось распознать документ из-за технической ошибки. Попробуйте ещё раз или отправьте более чёткое фото."
        )
        return
    except SpreadsheetInitError as exc:
        logger.exception(
            "[doc_id=%s user_id=%s] Spreadsheet initialization failed for file name=%s mime=%s",
            internal_doc_id,
            user_id,
            getattr(document, "file_name", None),
            getattr(document, "mime_type", None),
            exc_info=exc,
        )
        await message.answer(SPREADSHEET_ERROR_MESSAGE)
        return
    except Exception as exc:  # pragma: no cover - runtime errors
        error_code = internal_doc_id.split("-")[-1][:6]
        logger.exception(
            "[doc_id=%s user_id=%s] Unexpected error code=%s for file name=%s mime=%s: %s",
            internal_doc_id,
            user_id,
            error_code,
            getattr(document, "file_name", None),
            getattr(document, "mime_type", None),
            exc,
        )
        await message.answer(
            "Не удалось распознать документ из-за технической ошибки. Попробуйте ещё раз или отправьте более чёткое фото."
        )
        return

    if pending is None or doc_data is None:
        await message.answer(
            "Произошла техническая ошибка при подготовке документа. Попробуйте позже или обратитесь в поддержку."
        )
        return

    if duplicate_info:
        pending.duplicate_order_number = duplicate_info.order_number
        pending.duplicate_file_id = duplicate_info.file_id
        PENDING_UPLOADS[pending.internal_doc_id] = pending
        await message.answer(
            _build_duplicate_message(doc_data, duplicate_info.order_number),
            reply_markup=_build_duplicate_keyboard(pending.internal_doc_id),
        )
        return

    try:
        await _persist_document(pending, message)
    except PersistDriveError as exc:
        logger.exception(
            "[doc_id=%s user_id=%s] Failed to persist due to Drive error: %s",
            internal_doc_id,
            user_id,
            exc,
        )
        if str(exc) == "drive_quota":
            await message.answer(DRIVE_QUOTA_ERROR_MESSAGE)
        else:
            await message.answer(
                "Не удалось сохранить документ в Google Диск. Попробуйте позже или обратитесь к администратору."
            )


def _build_drive_filename(order_number: int, doc: ParsedDocument, suffix: str) -> str:
    counterparty = _transliterate_counterparty(doc.main_counterparty or "БезКонтрагента")
    doc_date = _format_date_for_filename(doc.date)
    extension = suffix.lower() if suffix else ".pdf"
    return f"{order_number}_{doc_date}_{counterparty}{extension}"


def _build_layout_copy_filename(order_number: int, doc: ParsedDocument, suffix: str) -> str:
    base = _build_drive_filename(order_number, doc, suffix)
    return base.replace(suffix, f"_copy{suffix}")


def _build_layout_title(doc: ParsedDocument) -> str:
    doc_date = doc.date or doc.doc_date
    date_part = doc_date.strftime("%d.%m.%Y") if doc_date else "—"
    number_part = (doc.number or doc.doc_number or "—").strip()
    type_part = doc.doc_type or "Документ"
    return f"{type_part} копия № {number_part} от {date_part}"


def _build_summary(doc: ParsedDocument, order_number: int, file_saved: bool) -> str:
    document_number = (doc.number or doc.doc_number or "").strip() or "—"
    doc_date = doc.date or doc.doc_date
    date_display = doc_date.strftime("%d.%m.%Y") if doc_date else "—"

    resolved_parties = resolve_parties(doc.supplier_raw, doc.customer_raw)

    def _display_name(value: str | None) -> str | None:
        if not value:
            return None
        return strip_address(value)

    supplier_display = _display_name(
        resolved_parties.executor_name
        or resolved_parties.executor_raw
        or (doc.supplier.name if doc.supplier else None)
        or doc.supplier_raw
    )
    customer_display = _display_name(
        resolved_parties.buyer_name
        or resolved_parties.buyer_raw
        or (doc.buyer.name if doc.buyer else None)
        or doc.customer_raw
    )

    currency = doc.currency if doc.currency not in {None, "UNKNOWN"} else None
    if not currency and doc.currency_code not in {None, "UNKNOWN"}:
        currency = doc.currency_code
    display_totals = doc.llm_totals or doc.totals or Totals()
    data_totals = doc.totals or display_totals
    currency = currency or display_totals.currency or data_totals.currency
    totals = getattr(doc, "totals", None)

    vat_rate = (
        display_totals.vat_rate_percent
        or display_totals.vat_rate_percent_text
        or data_totals.vat_rate_percent
        or data_totals.vat_rate_percent_text
    )

    def _vat_with_rate(
        value: float | None,
        currency: str | None,
        vat_amount_text: str | None = None,
        vat_rate_text: str | None = None,
    ) -> str:
        if value is None:
            return "—"

        base = _format_money(value, currency, text=vat_amount_text)
        if vat_rate_text:
            return f"{base} ({vat_rate_text})"
        return base

    vat_value = display_totals.vat_amount or display_totals.total_vat
    vat_amount_text = getattr(totals, "vat_amount_text", None) if totals else getattr(display_totals, "vat_amount_text", None)
    vat_rate_text = getattr(totals, "vat_rate_text", None) if totals else getattr(display_totals, "vat_rate_text", None)

    totals_block = [
        f"• Без НДС: {esc(_format_money(display_totals.total_without_vat or display_totals.amount_without_vat, currency, text=display_totals.total_without_vat_text))}",
        f"• Сумма НДС: {esc(_vat_with_rate(vat_value, currency, vat_amount_text, vat_rate_text))}",
        f"• С НДС: {esc(_format_money(display_totals.total_with_vat or display_totals.amount_with_vat, currency, text=display_totals.total_with_vat_text))}",
    ]
    if vat_rate:
        rate_text = format_money_ru(vat_rate, empty="") or str(vat_rate)
        totals_block.append(f"• Ставка НДС: {esc(rate_text)} %")

    warnings = list(doc.warnings)
    if doc.sums_suspect:
        warnings.append("⚠️ Итоговые суммы выглядят подозрительно и требуют проверки.")
    if doc.normalization_notes:
        warnings.extend(doc.normalization_notes)
    if not warnings and not doc.validation_errors():
        warning_line = "\n\n<b>Нужна проверка:</b>\n• Нет критичных предупреждений"
    else:
        warning_line = ""

    normalization_line = ""
    verification_lines: list[str] = []
    raw_supplier = doc.supplier_raw or resolved_parties.executor_raw or "—"
    raw_customer = doc.customer_raw or resolved_parties.buyer_raw or "—"

    supplier_name = resolved_parties.executor_name or doc.supplier_company_name
    customer_name = resolved_parties.buyer_name or doc.customer_company_name

    if supplier_name:
        verification_lines.append(
            f"• Исполнитель: {esc(raw_supplier)} → {esc(supplier_name)}"
            + (f" (код: {esc(resolved_parties.executor_id or doc.supplier_company_code or 'словарь')})" if resolved_parties.executor_in_dict or doc.supplier_company_code else "")
        )
    else:
        verification_lines.append(
            f"• Исполнитель: прочитано — {esc(raw_supplier)}; в словаре не найдено"
        )

    if customer_name:
        verification_lines.append(
            f"• Контрагент: {esc(raw_customer)} → {esc(customer_name)}"
            + (f" (код: {esc(resolved_parties.buyer_id or doc.customer_company_code or 'словарь')})" if resolved_parties.buyer_in_dict or doc.customer_company_code else "")
        )
    else:
        verification_lines.append(
            f"• Контрагент: прочитано — {esc(raw_customer)}; в словаре не найдено"
        )

    if doc.normalization_notes:
        normalized_notes = [note.replace("normalize:", "").strip() for note in doc.normalization_notes]
        verification_lines.extend(f"• {esc(note)}" for note in normalized_notes)
    if verification_lines:
        normalization_line = "\n\n<b>Проверка названий:</b>\n" + "\n".join(verification_lines)
    if warnings:
        warning_list = "\n".join(f"• {esc(w)}" for w in warnings)
        warning_line = f"\n\n<b>Нужна проверка:</b>\n{warning_list}"

    drive_line = "Файл сохранён на Google Диск." if file_saved else "Файл не сохранён на Google Диск."

    return (
        "Документ обработан.\n\n"
        f"Тип: <b>{esc(doc.doc_type or '—')}</b>\n"
        f"Номер: {esc(document_number)}\n"
        f"Дата: {esc(date_display)}\n"
        f"Контрагент: {esc(customer_display or 'не распознан')}\n"
        f"Исполнитель: {esc(supplier_display or 'не распознан')}\n"
        f"Валюта: {esc(currency or doc.currency_code or '—')}\n\n"
        f"<b>Итого по документу:</b>\n" + "\n".join(totals_block) + "\n\n"
        + f"Номер по порядку: {order_number}\n"
        f"{esc(drive_line)}{normalization_line}{warning_line}"
    )

def _build_duplicate_message(doc: ParsedDocument, existing_order: int) -> str:
    number = (doc.number or "").strip() or "—"
    date_display = doc.date.strftime("%d-%m-%Y") if doc.date else "—"
    counterparty = doc.main_counterparty or "—"
    return (
        "Документ с такими реквизитами уже есть в таблице.\n"
        f"Тип: {doc.doc_type}\n"
        f"Номер: {number}\n"
        f"Дата: {date_display}\n"
        f"Контрагент: {counterparty}\n"
        f"Найденная запись имеет номер по порядку: {existing_order}.\n"
        "Заменить её новым документом?"
    )


def _build_duplicate_keyboard(pending_id: str) -> InlineKeyboardMarkup:
    builder = InlineKeyboardBuilder()
    builder.button(text="🔄 Заменить запись", callback_data=f"dup:replace:{pending_id}")
    builder.button(text="↩️ Оставить как есть", callback_data=f"dup:keep:{pending_id}")
    builder.adjust(1)
    return builder.as_markup()


def _format_date_for_filename(value: date | None) -> str:
    target_date = value or date.today()
    return target_date.strftime("%d-%m-%Y")


def _transliterate_counterparty(name: str) -> str:
    transliteration_map = {
        "а": "a",
        "б": "b",
        "в": "v",
        "г": "g",
        "д": "d",
        "е": "e",
        "ё": "e",
        "ж": "zh",
        "з": "z",
        "и": "i",
        "й": "y",
        "к": "k",
        "л": "l",
        "м": "m",
        "н": "n",
        "о": "o",
        "п": "p",
        "р": "r",
        "с": "s",
        "т": "t",
        "у": "u",
        "ф": "f",
        "х": "h",
        "ц": "c",
        "ч": "ch",
        "ш": "sh",
        "щ": "sch",
        "ъ": "",
        "ы": "y",
        "ь": "",
        "э": "e",
        "ю": "yu",
        "я": "ya",
    }
    normalized = unicodedata.normalize("NFKD", name).lower()
    transliterated = "".join(transliteration_map.get(ch, ch) for ch in normalized)
    transliterated = re.sub(r"[^a-z0-9_-]+", "_", transliterated)
    transliterated = re.sub(r"_+", "_", transliterated).strip("_")
    return transliterated or "Bez_kontragenta"


def _extract_first_page_image(path: Path) -> bytes:
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            raise ValueError("PDF has no pages")
        page = pdf.pages[0]
        image = page.to_image(resolution=300)
        buffer = io.BytesIO()
        image.original.save(buffer, format="PNG")
        return buffer.getvalue()


def _extract_docx_text(path: Path) -> str:
    doc = Document(path)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            lines.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _postprocess_document(doc: ParsedDocument, our_org: str) -> None:
    warnings = list(doc.warnings)
    if doc.supplier and not doc.contractor:
        doc.contractor = doc.supplier.name
    if doc.buyer and not doc.executor:
        doc.executor = doc.buyer.name
    if doc.contractor and not doc.main_counterparty:
        doc.main_counterparty = doc.contractor
    if doc.contractor and not doc.counterparty:
        doc.counterparty = doc.contractor
    if doc.doc_number and not doc.number:
        doc.number = doc.doc_number
    if doc.doc_date and not doc.date:
        doc.date = doc.doc_date
    if doc.executor and not doc.supplier:
        doc.supplier = Party(name=doc.executor)

    detected_code = detect_currency(doc.raw_text or "", doc.currency)
    if detected_code:
        doc.currency_code = detected_code
        doc.currency = detected_code
    if not doc.currency_code:
        doc.currency_code = "UNKNOWN"
    if doc.currency_code == "UNKNOWN":
        warnings.append("Не удалось определить валюту автоматически")
    if doc.date and not (date(2010, 1, 1) <= doc.date <= date(2035, 12, 31)):
        warnings.append("Дата документа вне допустимого диапазона, требуется проверка")
    if not doc.date:
        warnings.append("Не удалось определить дату документа")
    if not doc.number:
        warnings.append("Не распознан номер документа")
    if not doc.main_counterparty:
        warnings.append("Не найден контрагент")
    if not doc.main_item and doc.items:
        doc.main_item = doc.items[0]
    if not doc.main_counterparty:
        doc.main_counterparty = _choose_counterparty(doc, our_org)
    doc.warnings = warnings


def _choose_counterparty(doc: ParsedDocument, our_org: str) -> str | None:
    our = our_org.lower().strip()
    buyer = (doc.buyer.name or "").lower() if doc.buyer else ""
    supplier = (doc.supplier.name or "").lower() if doc.supplier else ""

    def _is_our(value: str) -> bool:
        return bool(value) and (our and our in value)

    if doc.doc_type in {"акт", "счет-фактура", "счет-протокол", "комбинированный"}:
        if buyer and not _is_our(buyer):
            return doc.buyer.name  # type: ignore[union-attr]
    if doc.doc_type in {"товарная_накладная", "товарно-транспортная_накладная", "накладная"}:
        if supplier and not _is_our(supplier):
            return doc.supplier.name  # type: ignore[union-attr]
    if doc.buyer and doc.buyer.name and not _is_our(buyer):
        return doc.buyer.name
    if doc.supplier and doc.supplier.name and not _is_our(supplier):
        return doc.supplier.name
    return doc.contractor or doc.counterparty


def _format_money(value, currency: str | None, *, text: str | None = None) -> str:
    display = prefer_text_amount(text, value)
    if not display:
        return "—"
    currency_display = f" {currency}" if currency else ""
    return f"{display}{currency_display}"


__all__ = ["router"]
PROJECT_ROOT = Path(__file__).resolve().parent.parent

