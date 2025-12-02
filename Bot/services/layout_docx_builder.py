from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from services.layout_types import WordBox


@dataclass
class LayoutLine:
    words: list[WordBox]

    @property
    def y(self) -> float:
        return sum(w.y for w in self.words) / len(self.words)

    @property
    def height(self) -> float:
        return sum(w.h for w in self.words) / len(self.words)

    @property
    def text(self) -> str:
        return " ".join(word.text for word in sorted(self.words, key=lambda w: w.x))


@dataclass
class ParagraphSection:
    lines: list[LayoutLine]


@dataclass
class TableSection:
    lines: list[LayoutLine]
    column_centers: list[float]

    @property
    def columns(self) -> int:
        return len(self.column_centers)


def build_layout_docx(words: list[WordBox], clean_text: str, output_path: Path, *, title: str | None = None) -> None:
    doc = Document()
    if title:
        heading = doc.add_heading(title, level=0)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    lines = _group_lines(words)
    sections = _split_sections(lines)

    for section in sections:
        if isinstance(section, TableSection):
            _render_table(doc, section)
        else:
            for line in section.lines:
                doc.add_paragraph(line.text)
            doc.add_paragraph("")

    if clean_text:
        doc.add_paragraph()
        doc.add_heading("Очистка текста (GPT)", level=1)
        doc.add_paragraph(clean_text)

    doc.save(output_path)


def _group_lines(words: list[WordBox]) -> list[LayoutLine]:
    sorted_words = sorted(words, key=lambda w: (w.y, w.x))
    lines: list[LayoutLine] = []
    current: list[WordBox] = []
    if not sorted_words:
        return lines

    baseline_height = sorted_words[0].h or 1
    tolerance = max(6, baseline_height * 0.6)

    for word in sorted_words:
        if not current:
            current.append(word)
            continue
        if abs(word.y - current[-1].y) <= tolerance:
            current.append(word)
        else:
            lines.append(LayoutLine(words=list(current)))
            current = [word]
    if current:
        lines.append(LayoutLine(words=list(current)))
    return lines


def _split_sections(lines: list[LayoutLine]) -> list[object]:
    sections: list[object] = []
    current_para: list[LayoutLine] = []
    table_buffer: list[tuple[LayoutLine, list[float]]] = []
    previous_line: LayoutLine | None = None
    gap_threshold = _estimate_gap(lines)

    for line in lines:
        columns = _estimate_columns(line)
        if columns and (not table_buffer or _columns_compatible(columns, table_buffer[-1][1])):
            if current_para:
                sections.append(ParagraphSection(lines=list(current_para)))
                current_para.clear()
            table_buffer.append((line, columns))
        else:
            if table_buffer:
                sections.append(_table_from_buffer(table_buffer))
                table_buffer.clear()
            if previous_line and abs(line.y - previous_line.y) > gap_threshold:
                if current_para:
                    sections.append(ParagraphSection(lines=list(current_para)))
                    current_para.clear()
            current_para.append(line)
        previous_line = line

    if table_buffer:
        sections.append(_table_from_buffer(table_buffer))
    if current_para:
        sections.append(ParagraphSection(lines=list(current_para)))
    return sections


def _estimate_gap(lines: list[LayoutLine]) -> float:
    heights = [line.height for line in lines] or [10]
    median_height = sorted(heights)[len(heights) // 2]
    return median_height * 1.6


def _estimate_columns(line: LayoutLine) -> list[float] | None:
    if len(line.words) < 3:
        return None
    centers = [w.x + w.w / 2 for w in line.words]
    centers.sort()
    buckets: list[list[float]] = []
    tolerance = 32
    for center in centers:
        placed = False
        for bucket in buckets:
            if abs(bucket[-1] - center) <= tolerance:
                bucket.append(center)
                placed = True
                break
        if not placed:
            buckets.append([center])
    averaged = [sum(bucket) / len(bucket) for bucket in buckets if bucket]
    if len(averaged) < 2:
        return None
    return averaged


def _columns_compatible(current: list[float], previous: list[float]) -> bool:
    if abs(len(current) - len(previous)) > 1:
        return False
    tolerance = 50
    for idx, center in enumerate(current[: len(previous)]):
        if abs(center - previous[idx]) > tolerance:
            return False
    return True


def _table_from_buffer(buffer: list[tuple[LayoutLine, list[float]]]) -> TableSection:
    column_sets = [cols for _, cols in buffer if cols]
    max_len = max(len(c) for c in column_sets)
    merged: list[list[float]] = [[] for _ in range(max_len)]
    for cols in column_sets:
        for idx, center in enumerate(cols):
            merged[idx].append(center)
    averaged = [sum(col) / len(col) for col in merged if col]
    return TableSection(lines=[line for line, _ in buffer], column_centers=averaged)


def _render_table(doc: Document, section: TableSection) -> None:
    rows = len(section.lines)
    cols = max(1, section.columns)
    table = doc.add_table(rows=rows, cols=cols)
    for row_idx, line in enumerate(section.lines):
        cells = table.rows[row_idx].cells
        cell_texts = [""] * cols
        for word in sorted(line.words, key=lambda w: w.x):
            idx = _closest_column(word, section.column_centers)
            cell_texts[idx] = (cell_texts[idx] + " " + word.text).strip()
        for col_idx, value in enumerate(cell_texts):
            cells[col_idx].text = value
    doc.add_paragraph("")


def _closest_column(word: WordBox, centers: list[float]) -> int:
    center = word.x + word.w / 2
    distances = [abs(center - col) for col in centers]
    return distances.index(min(distances)) if distances else 0


__all__ = ["build_layout_docx", "LayoutLine", "TableSection", "ParagraphSection"]
