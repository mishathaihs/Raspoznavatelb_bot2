from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


def _write_header(doc: Document, title: str) -> None:
    heading = doc.add_heading(level=0)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = heading.add_run(title)
    run.font.size = Pt(16)
    run.bold = True


def _write_separator(doc: Document) -> None:
    paragraph = doc.add_paragraph("―" * 40)
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


def _add_info_paragraphs(doc: Document) -> None:
    doc.add_paragraph("Исполнитель: {{ executor_name }}")
    doc.add_paragraph("Заказчик: {{ customer_name }}")
    doc.add_paragraph("Дата: {{ doc_date }}")
    doc.add_paragraph("Номер: {{ doc_number }}")


def _add_info_table_v2(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Исполнитель"
    table.rows[0].cells[1].text = "{{ EXECUTOR }}"
    table.rows[1].cells[0].text = "Заказчик"
    table.rows[1].cells[1].text = "{{ CUSTOMER }}"
    doc.add_paragraph()


def _add_header_meta_v2(doc: Document) -> None:
    meta = doc.add_table(rows=1, cols=2)
    meta.rows[0].cells[0].text = "Номер: {{ DOC_NUMBER }}"
    meta.rows[0].cells[1].text = "Дата: {{ DOC_DATE }}"
    doc.add_paragraph()


def _add_table(doc: Document, is_waybill: bool) -> None:
    headers = ["Наименование", "Сумма без НДС", "Сумма НДС", "Сумма с НДС", "Валюта"]
    table = doc.add_table(rows=1, cols=len(headers))
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
    row = table.add_row().cells
    row[0].text = "{{ position_name }}"
    row[1].text = ""
    row[2].text = "{{ amount_vat }}"
    row[3].text = "{{ amount_total }}"
    row[4].text = "{{ currency }}"
    if is_waybill:
        doc.add_paragraph("Грузоотправитель/Поставщик: {{ executor_name }}")
        doc.add_paragraph("Грузополучатель/Покупатель: {{ customer_name }}")


def _add_table_v2(doc: Document) -> None:
    headers = [
        "Наименование",
        "Сумма без НДС",
        "Сумма НДС",
        "Сумма с НДС",
        "Валюта",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
    sample = table.add_row().cells
    sample[0].text = "{{ POSITION_NAME }}"
    sample[1].text = "{{ POSITION_AMOUNT_WO_VAT }}"
    sample[2].text = "{{ POSITION_VAT }}"
    sample[3].text = "{{ POSITION_WITH_VAT }}"
    sample[4].text = "{{ POSITION_CURRENCY }}"


def _add_totals(doc: Document) -> None:
    doc.add_paragraph("Итого с НДС: {{ amount_total }} {{ currency }}")
    doc.add_paragraph("В том числе НДС: {{ amount_vat }} {{ currency }}")


def _add_totals_v2(doc: Document) -> None:
    doc.add_paragraph("Итого без НДС: {{ TOTAL_WITHOUT_VAT }} {{ CURRENCY }}")
    doc.add_paragraph("НДС: {{ TOTAL_VAT }} {{ CURRENCY }}")
    doc.add_paragraph("Итого с НДС: {{ TOTAL_WITH_VAT }} {{ CURRENCY }}")


def create_act_template(path: Path) -> None:
    doc = Document()
    _write_header(doc, "АКТ оказанных услуг")
    _add_info_paragraphs(doc)
    _add_table(doc, is_waybill=False)
    _add_totals(doc)
    doc.save(path)


def create_act_template_v2(path: Path) -> None:
    doc = Document()
    _write_header(doc, "АКТ оказанных услуг")
    _write_separator(doc)
    _add_header_meta_v2(doc)
    _add_info_table_v2(doc)
    _add_table_v2(doc)
    _add_totals_v2(doc)
    doc.save(path)


def create_waybill_template(path: Path) -> None:
    doc = Document()
    _write_header(doc, "ТОВАРНАЯ НАКЛАДНАЯ")
    _add_info_paragraphs(doc)
    _add_table(doc, is_waybill=True)
    _add_totals(doc)
    doc.save(path)


def main() -> None:
    templates_dir = Path("templates")
    templates_dir.mkdir(exist_ok=True)

    act_path = templates_dir / "template_act.docx"
    waybill_path = templates_dir / "template_waybill.docx"
    act_v2_path = templates_dir / "act_v2.docx"

    create_act_template(act_path)
    create_waybill_template(waybill_path)
    create_act_template_v2(act_v2_path)
    print(f"Generated templates: {act_path}, {waybill_path}, {act_v2_path}")


if __name__ == "__main__":
    main()
